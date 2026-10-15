"""Generate local, editable Portuguese petition drafts as text or Word.

The inherited contribution-time template is a starting point for a lawyer's
review, not a calculation or legal eligibility engine. Word support is an
optional dependency; data is never sent to an external service.
"""

from __future__ import annotations

import argparse
import datetime as _dt
from dataclasses import dataclass
from io import BytesIO
import os
from pathlib import Path
import re
import tempfile
from typing import Any, Dict, Iterable, List, TYPE_CHECKING

if TYPE_CHECKING:
    from docx.document import Document as WordDocument

MONTHS_PT = {
    1: "janeiro",
    2: "fevereiro",
    3: "março",
    4: "abril",
    5: "maio",
    6: "junho",
    7: "julho",
    8: "agosto",
    9: "setembro",
    10: "outubro",
    11: "novembro",
    12: "dezembro",
}


def load_case_data(source: Path) -> Dict[str, Any]:
    """Load the YAML case description from ``source``.

    This implementation intentionally avoids external dependencies so the
    project can run in restricted environments.  The supported YAML subset
    covers dictionaries, lists and scalar values (strings, numbers and
    booleans), which is sufficient for the structured data used by the
    application.
    """

    try:
        text = source.read_text(encoding="utf-8")
    except FileNotFoundError as exc:  # pragma: no cover - thin wrapper
        raise SystemExit("Arquivo de entrada não encontrado.") from exc

    return parse_simple_yaml(text)


def parse_simple_yaml(text: str) -> Dict[str, Any]:
    """Parse a minimal subset of YAML into Python structures."""

    lines = [
        (len(line) - len(line.lstrip(" ")), line.strip())
        for line in text.splitlines()
        if line.strip() and not line.strip().startswith("#")
    ]

    value, index = _parse_block(lines, 0, 0)
    if index != len(lines):  # pragma: no cover - defensive branch
        raise ValueError("Não foi possível interpretar todo o arquivo YAML.")
    if not isinstance(value, dict):
        raise ValueError("A raiz do YAML deve ser um objeto (mapping).")
    return value


def _parse_block(lines: List[tuple[int, str]], start: int, indent: int):
    result: Any = None
    idx = start

    while idx < len(lines):
        line_indent, content = lines[idx]
        if line_indent < indent:
            break

        if content.startswith("- "):
            if result is None:
                result = []
            elif not isinstance(result, list):
                break

            value_part = content[2:].strip()
            idx += 1

            if not value_part:
                item, idx = _parse_block(lines, idx, line_indent + 2)
                result.append(item)
                continue

            if ":" in value_part:
                key, raw_val = value_part.split(":", 1)
                key = key.strip()
                raw_val = raw_val.strip()
                item: Dict[str, Any] = {}

                if raw_val:
                    item[key] = _parse_scalar(raw_val)
                else:
                    child, idx = _parse_block(lines, idx, line_indent + 2)
                    item[key] = child

                if idx < len(lines) and lines[idx][0] > line_indent:
                    extra, idx = _parse_block(lines, idx, line_indent + 2)
                    if isinstance(extra, dict):
                        item.update(extra)
                    else:  # pragma: no cover - malformed structure
                        raise ValueError("Estrutura de lista inválida.")

                result.append(item)
            else:
                result.append(_parse_scalar(value_part))
                if idx < len(lines) and lines[idx][0] > line_indent:
                    extra, idx = _parse_block(lines, idx, line_indent + 2)
                    result[-1] = extra

            continue

        if result is None:
            result = {}
        elif not isinstance(result, dict):
            break

        if ":" not in content:
            # Never include a client's input line in error messages/logs.
            raise ValueError("Linha inválida no YAML.")

        key, raw_val = content.split(":", 1)
        key = key.strip()
        raw_val = raw_val.strip()
        idx += 1

        if raw_val:
            result[key] = _parse_scalar(raw_val)
        else:
            child, idx = _parse_block(lines, idx, line_indent + 2)
            result[key] = child

    if result is None:
        result = {}

    return result, idx


def _parse_scalar(value: str) -> Any:
    lowered = value.lower()
    if lowered in {"true", "yes"}:
        return True
    if lowered in {"false", "no"}:
        return False
    if lowered in {"null", "none", "~"}:
        return None

    if value.startswith(("'", '"')) and value.endswith(("'", '"')):
        return value[1:-1]

    try:
        if "." in value:
            return float(value)
        return int(value)
    except ValueError:
        return value


def format_date(date_str: str) -> str:
    """Convert an ISO date string (YYYY-MM-DD) to long Portuguese format."""

    if not date_str:
        return ""

    try:
        if not isinstance(date_str, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", date_str):
            raise ValueError
        date = _dt.date.fromisoformat(date_str)
    except ValueError:
        raise ValueError("Data inválida; use uma data real no formato AAAA-MM-DD.") from None
    return f"{date.day:02d} de {MONTHS_PT[date.month]} de {date.year}"


def normalize_oab(value: Any) -> str:
    """Use one OAB prefix, preserving the supplied registration and its zeros.

    Accept e.g. ``PR 00000``, ``OAB/PR 00000`` and legacy ``OAB OAB/PR 00000``.
    This is display normalization, not verification of a professional record.
    """
    text = "" if value is None else " ".join(str(value).split())
    text = re.sub(r"^(?:OAB\b\s*[/:-]?\s*)+", "", text, flags=re.IGNORECASE)
    if not text:
        return "OAB ***"
    state = re.match(r"^([a-zA-Z]{2})(?=\s|\d|/|-|$)\s*[/\-]?\s*(.*)$", text)
    if state:
        return f"OAB/{state[1].upper()} {state[2]}".rstrip()
    return f"OAB {text}"


def _format_vinculos(vinculos: Iterable[Dict[str, Any]]) -> List[str]:
    formatted = []
    for vinculo in vinculos:
        inicio = format_date(vinculo.get("inicio", ""))
        fim = format_date(vinculo.get("fim", ""))
        categoria = vinculo.get("categoria", "não informado")
        formatted.append(
            f"Período de {inicio} a {fim}, na qualidade de segurada {categoria}."
        )
    return formatted


def _format_especial(especial: Iterable[Dict[str, Any]]) -> List[str]:
    formatted = []
    for periodo in especial:
        inicio = format_date(periodo.get("inicio", ""))
        fim = format_date(periodo.get("fim", ""))
        agente = periodo.get("agente", "agente nocivo não informado")
        ppp = "com PPP anexado" if periodo.get("ppp") else "sem PPP comprovado"
        formatted.append(
            f"Atividade especial de {inicio} a {fim}, exposta a {agente}, {ppp}."
        )
    return formatted


@dataclass(frozen=True)
class _Paragraph:
    text: str
    style: str = "Normal"


def _validate_case_data(case_data: Dict[str, Any]) -> None:
    """Check structure without putting private field values into exceptions."""
    if not isinstance(case_data, dict):
        raise ValueError("Os dados do caso devem ser um objeto (mapping).")
    for name in ("processo", "partes", "advogada", "beneficio", "tempos", "pedidos"):
        if name in case_data and not isinstance(case_data[name], dict):
            raise ValueError(f"O campo {name} deve ser um objeto (mapping).")
    tempos = case_data.get("tempos", {})
    for name in ("vinculos", "especial"):
        values = tempos.get(name, [])
        if not isinstance(values, list) or any(not isinstance(v, dict) for v in values):
            raise ValueError(f"O campo tempos.{name} deve ser uma lista de objetos.")
    if case_data.get("observacoes") is not None and not isinstance(case_data["observacoes"], str):
        raise ValueError("O campo observacoes deve ser um texto.")
    if not isinstance(case_data.get("beneficio", {}).get("tipo", ""), str):
        raise ValueError("O campo beneficio.tipo deve ser um texto.")


def _petition_paragraphs(case_data: Dict[str, Any]) -> List[_Paragraph]:
    """Build content once so Word and plain text cannot silently diverge."""
    _validate_case_data(case_data)
    processo = case_data.get("processo", {})
    partes = case_data.get("partes", {})
    advogada = case_data.get("advogada", {})
    beneficio = case_data.get("beneficio", {})
    tempos = case_data.get("tempos", {})
    pedidos = case_data.get("pedidos", {})
    observacoes = case_data.get("observacoes")
    oab = normalize_oab(advogada.get("oab"))

    paragraphs = [
        _Paragraph((
            "Excelentíssimo(a) Senhor(a) Doutor(a) Juiz(a) Federal da "
            f"{processo.get('foro', 'Vara Federal')} da Seção Judiciária "
            f"{processo.get('secao', '')}"
        ).strip(), "Enderecamento"),
        _Paragraph(
            f"{partes.get('autora', 'A autora')}, brasileira, portadora do CPF nº "
            f"{partes.get('cpf', '***')}, residente na {partes.get('endereco', '')}, "
            f"por sua advogada que esta subscreve ({oab}), com endereço eletrônico "
            f"{advogada.get('email', '***')}, vem, respeitosamente, propor a presente"
        ),
        _Paragraph(
            "AÇÃO PREVIDENCIÁRIA DE CONCESSÃO DE BENEFÍCIO por tempo de contribuição.",
            "Title",
        ),
        _Paragraph("1. DOS FATOS", "Heading 1"),
        _Paragraph(
            f"Requer a concessão do benefício de {beneficio.get('tipo', '').upper()}"
            f", com DER em {format_date(beneficio.get('der', ''))}, sob a regra de "
            f"transição {beneficio.get('regra_transicao', 'não informada')}."
        ),
        _Paragraph("1.1. Do tempo rural/urbano", "Heading 2"),
    ]
    vinculos = _format_vinculos(tempos.get("vinculos", []))
    paragraphs.extend(_Paragraph(v) for v in (vinculos or ["Não foram informados vínculos comuns."]))
    paragraphs.append(_Paragraph("1.2. Do tempo especial", "Heading 2"))
    especial = _format_especial(tempos.get("especial", []))
    paragraphs.extend(_Paragraph(v) for v in (especial or ["Não foram informados períodos especiais."]))
    if tempos.get("conversao_especial"):
        paragraphs.append(_Paragraph(
            "Requer a conversão do tempo especial em comum, com aplicação do fator correspondente."
        ))
    carencia = case_data.get("carencia_meses")
    paragraphs.extend([
        _Paragraph("1.3. Da carência", "Heading 2"),
        _Paragraph(
            f"Foi informado um total de {carencia} contribuições, sujeito à conferência "
            "da carência aplicável."
            if carencia is not None else "A carência deverá ser comprovada em instrução."
        ),
        _Paragraph("2. DOS PEDIDOS", "Heading 1"),
    ])
    requests = []
    if pedidos.get("gratuidade"):
        requests.append("A concessão dos benefícios da justiça gratuita.")
    requests.append("O reconhecimento dos períodos de labor acima descritos, com a devida averbação.")
    if pedidos.get("tutela_implantacao"):
        requests.append("A concessão de tutela de urgência para imediata implantação do benefício.")
    juros = pedidos.get("correcao_juros", "")
    if juros:
        requests.append(f"A aplicação do índice de atualização {juros} até o efetivo pagamento.")
    requests.append(
        "A condenação do INSS ao pagamento das parcelas vencidas, acrescidas de juros e correção monetária."
    )
    for index, request in enumerate(requests):
        paragraphs.append(_Paragraph(f"{chr(97 + index)}) {request}", "Pedido"))
    if observacoes:
        paragraphs.extend([
            _Paragraph("3. DAS PROVAS", "Heading 1"),
            _Paragraph("A parte autora se vale dos documentos já acostados, em especial:"),
        ])
        # Input remains plain text; it is never interpreted as XML or formatting.
        paragraphs.extend(_Paragraph(line) for line in observacoes.splitlines() if line.strip())

    data_assinatura = format_date(case_data.get("data_assinatura") or _dt.date.today().isoformat())
    paragraphs.extend([
        _Paragraph("Termos em que,", "Fecho"),
        _Paragraph("Pede deferimento.", "Fecho"),
        _Paragraph(f"{processo.get('secao', '')}, {data_assinatura}.", "Assinatura"),
        _Paragraph(str(advogada.get("nome", "Advogada")), "Assinatura"),
        _Paragraph(oab, "Assinatura"),
    ])
    return paragraphs


def generate_petition(case_data: Dict[str, Any]) -> str:
    """Generate an editable plain-text draft without third-party dependencies."""
    return "\n\n".join(p.text for p in _petition_paragraphs(case_data))


def generate_petition_docx(
    case_data: Dict[str, Any], *, page_size: str = "letter"
) -> "WordDocument":
    """Return a native, editable python-docx document; no network or logging.

    Use ``document.save(path)`` or ``save_petition`` to persist it. Only this
    function needs python-docx; the plain-text API stays dependency-free.
    """
    if page_size not in {"a4", "letter"}:
        raise ValueError("Tamanho de página inválido; use a4 ou letter.")
    paragraphs = _petition_paragraphs(case_data)
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Inches, Pt, RGBColor
    except ImportError:
        raise RuntimeError(
            "Para gerar Word, instale: python -m pip install -r requirements-petitions.txt"
        ) from None

    document = Document()
    section = document.sections[0]
    if page_size == "a4":
        section.page_width, section.page_height = Cm(21), Cm(29.7)
    else:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.left_margin = Cm(3)
    section.bottom_margin = section.right_margin = Cm(2)
    section.header_distance = section.footer_distance = Cm(1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(1.25)
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.widow_control = True
    pf.keep_with_next = False
    pf.keep_together = False
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "pt-BR")
    normal.element.get_or_add_rPr().append(lang)

    for name in ("Title", "Heading 1", "Heading 2", "Enderecamento", "Pedido", "Fecho", "Assinatura", "Footer"):
        style = document.styles[name] if name in document.styles else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name in {"Title", "Heading 1", "Heading 2", "Enderecamento"}
        style.font.italic = False
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.keep_with_next = name in {"Title", "Heading 1", "Heading 2", "Enderecamento"}
        style.paragraph_format.keep_together = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Built-in Word styles can carry theme fonts and a blue Title border.
    # Explicit font names alone do not override every theme attribute in Word/LO.
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Enderecamento", "Pedido", "Fecho", "Assinatura", "Footer"):
        style = document.styles[name]
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in list(fonts.attrib):
            if attribute.lower().endswith("theme"):
                del fonts.attrib[attribute]
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{script}"), "Times New Roman")
        style.font.underline = False
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)

    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].paragraph_format.space_before = Pt(12)
    document.styles["Title"].font.size = Pt(14)
    document.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.styles["Enderecamento"].paragraph_format.space_after = Pt(18)
    pedido = document.styles["Pedido"].paragraph_format
    pedido.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pedido.left_indent = Cm(0.75)
    pedido.first_line_indent = Cm(-0.75)
    pedido.keep_together = False
    for name in ("Fecho", "Assinatura"):
        document.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.styles[name].paragraph_format.line_spacing = 1.0
        document.styles[name].paragraph_format.space_after = Pt(3)
        document.styles[name].paragraph_format.keep_with_next = True

    for block in paragraphs:
        p = document.add_paragraph(block.text, style=block.style)
        if block.style == "Fecho" and block.text == "Termos em que,":
            p.paragraph_format.space_before = Pt(12)
    # Keep the closing and signature together, but never bind them to a next page.
    document.paragraphs[-1].paragraph_format.keep_with_next = False
    document.paragraphs[-3].paragraph_format.space_before = Pt(12)

    footer = section.footer.paragraphs[0]
    footer.style = document.styles["Footer"]
    footer.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.line_spacing = 1.0
    document.styles["Footer"].font.size = Pt(10)
    footer.add_run("Página ")
    for node_type in ("begin", "instruction", "separate", "value", "end"):
        run = footer.add_run()
        if node_type == "instruction":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = " PAGE "
        elif node_type == "value":
            element = OxmlElement("w:t")
            element.text = "1"
        else:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), node_type)
        run._r.append(element)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    document.settings.element.append(update)

    # Do not copy client data, local user names or paths into document metadata.
    props = document.core_properties
    for name in ("author", "last_modified_by", "subject", "keywords", "comments", "identifier", "category"):
        setattr(props, name, "")
    props.title = "Minuta de petição previdenciária"
    props.language = "pt-BR"
    return document


def _resolve_format(output: Path | None, output_format: str | None) -> str:
    suffix = output.suffix.lower() if output else ""
    chosen = output_format or ("docx" if suffix == ".docx" else "txt")
    if chosen not in {"txt", "docx"}:
        raise ValueError("Formato inválido; use txt ou docx.")
    if suffix and suffix not in {".txt", ".docx"}:
        raise ValueError("A extensão da saída deve ser .txt ou .docx.")
    if suffix and suffix != f".{chosen}":
        raise ValueError("A extensão da saída não corresponde ao formato solicitado.")
    if chosen == "docx" and output is None:
        raise ValueError("A saída Word exige um arquivo: use -o peticao.docx.")
    return chosen


def save_petition(
    case_data: Dict[str, Any], output: Path, *, output_format: str | None = None,
    page_size: str = "letter",
) -> None:
    """Write locally and atomically, with private permissions on POSIX systems."""
    output = Path(output)
    chosen = _resolve_format(output, output_format)
    if chosen == "docx":
        buffer = BytesIO()
        generate_petition_docx(case_data, page_size=page_size).save(buffer)
        payload = buffer.getvalue()
    else:
        payload = generate_petition(case_data).encode("utf-8")
    output.parent.mkdir(parents=True, exist_ok=True)
    # Atomic replacement avoids leaving a truncated petition after a failed write.
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=output.parent, prefix=".petition-", delete=False) as stream:
            temporary = Path(stream.name)
            stream.write(payload)
        os.replace(temporary, output)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def build_cli() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Gera minuta de petição previdenciária em texto ou Word a partir de um YAML.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Exemplos de uso:\n"
            "  python petition_generator.py -i cases/exemplo_ficticio.yaml -o out/peticao.txt\n"
            "  python petition_generator.py -i cases/exemplo_ficticio.yaml -o out/peticao.docx --page-size a4"
        ),
    )
    parser.add_argument("-i", "--input", required=True, type=Path, help="Arquivo YAML com os dados do processo.")
    parser.add_argument("-o", "--output", type=Path, help="Saída .txt ou .docx. Sem saída, imprime texto no terminal.")
    parser.add_argument("--format", choices=("txt", "docx"), help="Formato explícito; por padrão é inferido da extensão.")
    parser.add_argument("--page-size", choices=("a4", "letter"), default="letter", help="Papel do Word: a4 ou letter (padrão).")
    return parser


def main(args: List[str] | None = None) -> None:
    parser = build_cli()
    namespace = parser.parse_args(args)
    try:
        chosen = _resolve_format(namespace.output, namespace.format)
        if namespace.output and namespace.input.resolve() == namespace.output.resolve():
            raise ValueError("Entrada e saída devem ser arquivos diferentes.")
    except ValueError as exc:
        # These fixed option messages contain no case data or user paths.
        parser.error(str(exc))
    try:
        case_data = load_case_data(namespace.input)
        if namespace.output:
            save_petition(case_data, namespace.output, output_format=chosen, page_size=namespace.page_size)
        else:
            print(generate_petition(case_data))
    except (ValueError, TypeError, AttributeError):
        # Parser/type/encoding errors can contain private input; never echo them.
        parser.error("Dados ou opções inválidos. Confira o YAML, as datas e o formato/extensão da saída.")
    except OSError:
        parser.error("Não foi possível ler ou gravar os arquivos. Confira os caminhos e as permissões.")
    except RuntimeError as exc:
        parser.error(str(exc))


if __name__ == "__main__":  # pragma: no cover - CLI entry point
    main()
