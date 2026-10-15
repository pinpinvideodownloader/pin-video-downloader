"""Regression tests with synthetic data only; never use client fixtures."""
from __future__ import annotations

import builtins
from copy import deepcopy
from io import BytesIO
import os
from pathlib import Path
import socket
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from petition_generator import (
    format_date, generate_petition, generate_petition_docx, load_case_data,
    main, normalize_oab, parse_simple_yaml, save_petition,
)

EXAMPLE = Path(__file__).resolve().parents[1] / "cases" / "exemplo_ficticio.yaml"


@pytest.fixture
def case():
    return load_case_data(EXAMPLE)


@pytest.mark.parametrize("value, expected", [
    ("OAB/PR 00000", "OAB/PR 00000"),
    ("PR 00000", "OAB/PR 00000"),
    ("OAB PR 00000", "OAB/PR 00000"),
    ("OAB OAB/PR 00000", "OAB/PR 00000"),
    ("  oab / pr   00000  ", "OAB/PR 00000"),
    ("OAB/SC 00001-A", "OAB/SC 00001-A"),
    ("00000", "OAB 00000"),
    (None, "OAB ***"), ("", "OAB ***"), ("OAB", "OAB ***"),
])
def test_normalize_oab(value, expected):
    assert normalize_oab(value) == expected
    assert normalize_oab(expected) == expected


@pytest.mark.parametrize("oab", ["OAB/PR 00000", "PR 00000", "OAB OAB/PR 00000"])
def test_oab_in_both_qualification_and_signature(case, oab):
    case["advogada"]["oab"] = oab
    for text in (generate_petition(case), "\n".join(p.text for p in generate_petition_docx(case).paragraphs)):
        assert text.count("OAB/PR 00000") == 2
        assert "OAB OAB" not in text
        assert text.endswith("OAB/PR 00000")


def test_docx_round_trip_matches_plain_text_and_is_editable(case, tmp_path):
    case["observacoes"] = "Primeiro parágrafo: ação & revisão <fictícia>.\n\nSegundo parágrafo com ç e ã."
    target = tmp_path / "nested" / "peticao.docx"
    before = deepcopy(case)
    save_petition(case, target)
    document = Document(target)
    assert [p.text for p in document.paragraphs] == generate_petition(case).split("\n\n")
    assert case == before  # Formatting never mutates the input.
    document.paragraphs[1].add_run(" Revisão fictícia.")
    document.save(target)
    assert Document(target).paragraphs[1].text.endswith("Revisão fictícia.")
    assert not document.settings.element.xpath(".//w:documentProtection")
    assert not document.inline_shapes


def test_docx_legal_styles_and_pagination(case):
    document = generate_petition_docx(case, page_size="a4")
    section = document.sections[0]
    assert section.page_width.cm == pytest.approx(21, abs=.01)
    assert section.page_height.cm == pytest.approx(29.7, abs=.01)
    for value, expected in [(section.top_margin, 3), (section.left_margin, 3),
                            (section.right_margin, 2), (section.bottom_margin, 2)]:
        assert value.cm == pytest.approx(expected, abs=.01)
    normal = document.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size.pt == 12
    assert normal.paragraph_format.line_spacing == 1.5
    assert normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert normal.paragraph_format.first_line_indent.cm == pytest.approx(1.25, abs=.01)
    assert normal.paragraph_format.widow_control is True
    assert normal.element.xpath("./w:rPr/w:lang/@w:val") == ["pt-BR"]
    assert [p.text for p in document.paragraphs if p.style.name == "Heading 1"] == [
        "1. DOS FATOS", "2. DOS PEDIDOS", "3. DAS PROVAS",
    ]
    assert len([p for p in document.paragraphs if p.style.name == "Title"]) == 1
    for name in ("Title", "Heading 1", "Heading 2", "Enderecamento"):
        style = document.styles[name]
        assert str(style.font.color.rgb) == "000000"
        assert style.paragraph_format.keep_with_next
        assert not style.element.xpath("./w:pPr/w:pBdr")
        fonts = style.element.xpath("./w:rPr/w:rFonts")[0]
        assert not any(attr.lower().endswith("theme") for attr in fonts.attrib)
    assert document.paragraphs[-2].style.paragraph_format.keep_with_next
    assert document.paragraphs[-1].paragraph_format.keep_with_next is False
    footer = section.footer.paragraphs[0]
    assert footer._p.xpath(".//w:instrText/text()") == [" PAGE "]
    assert document.settings.element.xpath("./w:updateFields/@w:val") == ["true"]


def test_default_paper_and_invalid_paper(case):
    section = generate_petition_docx(case).sections[0]
    assert section.page_width.inches == 8.5
    assert section.page_height.inches == 11
    with pytest.raises(ValueError):
        generate_petition_docx(case, page_size="unknown")


def test_docx_metadata_and_no_external_relationships(case):
    document = generate_petition_docx(case)
    props = document.core_properties
    assert props.author == props.last_modified_by == props.comments == ""
    buffer = BytesIO()
    document.save(buffer)
    with ZipFile(buffer) as archive:
        metadata = archive.read("docProps/core.xml").decode()
        for private_value in [case["partes"]["autora"], case["partes"]["cpf"], case["advogada"]["email"]]:
            assert private_value not in metadata
        assert not any(name.endswith("vbaProject.bin") for name in archive.namelist())
        assert not any(b'TargetMode="External"' in archive.read(name)
                       for name in archive.namelist() if name.endswith(".rels"))


def test_no_network_during_generation(case, tmp_path, monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Petition generation attempted a network connection")
    monkeypatch.setattr(socket, "socket", forbidden)
    save_petition(case, tmp_path / "local.docx")
    save_petition(case, tmp_path / "local.txt")


def test_txt_remains_usable_without_python_docx(case, monkeypatch):
    real_import = builtins.__import__
    def limited_import(name, *args, **kwargs):
        if name == "docx" or name.startswith("docx."):
            raise ImportError("dependency intentionally unavailable")
        return real_import(name, *args, **kwargs)
    monkeypatch.setattr(builtins, "__import__", limited_import)
    assert "AÇÃO PREVIDENCIÁRIA" in generate_petition(case)
    with pytest.raises(RuntimeError, match="requirements-petitions.txt"):
        generate_petition_docx(case)


@pytest.mark.parametrize("options, suffix", [
    ([], ".docx"), ([], ".DOCX"), (["--format", "docx"], ".docx"),
    (["--format", "docx"], ""), ([], ".txt"), (["--format", "txt"], ".txt"),
])
def test_cli_output_modes(tmp_path, capsys, options, suffix):
    target = tmp_path / "new" / ("peticao" + suffix)
    main(["-i", str(EXAMPLE), "-o", str(target), *options])
    assert capsys.readouterr().out == ""
    if "docx" in options or suffix.lower() == ".docx":
        assert Document(target).paragraphs[0].text.startswith("Excelentíssimo")
    else:
        assert "AÇÃO PREVIDENCIÁRIA" in target.read_text(encoding="utf-8")


def test_cli_stdout_remains_plain_text(capsys):
    main(["-i", str(EXAMPLE)])
    result = capsys.readouterr()
    assert result.out == generate_petition(load_case_data(EXAMPLE)) + "\n"
    assert result.err == ""


@pytest.mark.parametrize("options", [
    ["--format", "docx"],
    ["--format", "txt", "-o", "result.docx"],
    ["--format", "docx", "-o", "result.txt"],
    ["-o", "result.pdf"],
])
def test_cli_rejects_incompatible_outputs(tmp_path, monkeypatch, capsys, options):
    monkeypatch.chdir(tmp_path)
    with pytest.raises(SystemExit) as exc:
        main(["-i", str(EXAMPLE), *options])
    assert exc.value.code == 2
    assert not list(tmp_path.iterdir())
    assert capsys.readouterr().out == ""


def test_cli_cannot_overwrite_its_input(tmp_path, capsys):
    path = tmp_path / "source"
    path.write_text(EXAMPLE.read_text(), encoding="utf-8")
    before = path.read_bytes()
    with pytest.raises(SystemExit):
        main(["-i", str(path), "-o", str(path)])
    assert path.read_bytes() == before
    assert not capsys.readouterr().out


@pytest.mark.parametrize("content", [
    "CONFIDENCIAL_FICTICIO_SEM_DOIS_PONTOS",
    'beneficio:\n  der: "CONFIDENCIAL_FICTICIO"',
    'partes: "CONFIDENCIAL_FICTICIO"',
    'observacoes: 123',
])
def test_invalid_input_does_not_echo_case_data(tmp_path, capsys, content):
    source = tmp_path / "case.yaml"
    source.write_text(content, encoding="utf-8")
    output = tmp_path / "case.docx"
    with pytest.raises(SystemExit) as exc:
        main(["-i", str(source), "-o", str(output)])
    result = capsys.readouterr()
    assert exc.value.code == 2
    assert "CONFIDENCIAL" not in result.err + result.out
    assert "Traceback" not in result.err
    assert not output.exists()


def test_yaml_parser_does_not_echo_invalid_line():
    with pytest.raises(ValueError) as exc:
        parse_simple_yaml("SEGREDO_FICTICIO")
    assert "SEGREDO_FICTICIO" not in str(exc.value)


@pytest.mark.parametrize("value", ["2026-02-30", "2026-13-01", "2026-1-01", "nonsense", 123])
def test_invalid_dates(value):
    with pytest.raises(ValueError, match="Data inválida"):
        format_date(value)


def test_empty_case_and_optional_sections():
    text = generate_petition({})
    assert "Não foram informados vínculos comuns." in text
    assert "Não foram informados períodos especiais." in text
    assert "3. DAS PROVAS" not in text
    assert "OAB ***" in text
    assert [p.text[:2] for p in generate_petition_docx({}).paragraphs if p.style.name == "Pedido"] == ["a)", "b)"]


def test_carencia_does_not_claim_eligibility(case):
    case["carencia_meses"] = 1
    text = generate_petition(case)
    assert "suficientes para a carência" not in text
    assert "1 contribuições, sujeito à conferência" in text


def test_fixed_signature_date_and_leap_day(case):
    assert "15 de outubro de 2026" in generate_petition(case)
    assert format_date("2024-02-29") == "29 de fevereiro de 2024"


def test_failed_write_preserves_previous_file_and_cleans_temporary(case, tmp_path, monkeypatch):
    output = tmp_path / "petition.docx"
    output.write_bytes(b"previous document")
    def fail(*args):
        raise OSError("synthetic failure")
    monkeypatch.setattr(os, "replace", fail)
    with pytest.raises(OSError):
        save_petition(case, output)
    assert output.read_bytes() == b"previous document"
    assert not list(tmp_path.glob(".petition-*"))


@pytest.mark.skipif(os.name != "posix", reason="POSIX file permissions")
def test_private_output_permissions(case, tmp_path):
    output = tmp_path / "petition.txt"
    save_petition(case, output)
    assert output.stat().st_mode & 0o777 == 0o600
